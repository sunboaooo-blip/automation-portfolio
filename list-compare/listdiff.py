#!/usr/bin/env python3
"""Compare two contact lists on ONE name column and report who is missing from which list.

usage: listdiff.py A.(csv|xlsx) B.(csv|xlsx) OUT.xlsx --col "Full Name" [--col-b "Name"]
                   [--norm exact|space|case|nfkc] [--sheet-a NAME] [--sheet-b NAME]

Normalisation levels (cumulative, chosen and stated before work; nothing looser is applied silently):
  exact  - byte-for-byte equal text
  space  - trim ends, collapse internal runs of whitespace to one space
  case   - space + case-insensitive (str.casefold)
  nfkc   - case + Unicode NFKC (full-width letters, ligatures, compatibility forms); accents are NOT removed
Per-record policy:
  - every source row is kept whole (all columns, original values) in the output;
  - a name counts as present in the other list if it occurs there at least once after normalisation;
  - blank names are never matched; they are listed separately;
  - names occurring more than once in a list are NOT collapsed: every row is reported, and the repeat is listed.
Output sheets: Missing from B (rows of A), Missing from A (rows of B), Blank names, Repeated names, Note.
Self-check: for each list, matched + missing + blank = total rows; otherwise exit 3 and OUT is deleted.
"""
import argparse
import csv
import os
import re
import sys
import unicodedata
from collections import Counter


def load(path, sheet=None):
    if path.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet] if sheet else wb.worksheets[0]
        rows = [["" if v is None else str(v) for v in r] for r in ws.iter_rows(values_only=True)]
    else:
        raw = open(path, "rb").read()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError as e:
            sys.exit(f"{path}: not UTF-8 ({e}); agree the encoding first")
        rows = list(csv.reader(text.splitlines(keepends=True)))
    if not rows:
        sys.exit(f"{path}: empty")
    return rows[0], rows[1:]


def normaliser(level):
    def f(s):
        if level in ("space", "case", "nfkc"):
            s = re.sub(r"\s+", " ", s).strip()
        if level == "nfkc":
            s = unicodedata.normalize("NFKC", s)
            s = re.sub(r"\s+", " ", s).strip()
        if level in ("case", "nfkc"):
            s = s.casefold()
        return s
    return f


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("a"); ap.add_argument("b"); ap.add_argument("out")
    ap.add_argument("--col", required=True); ap.add_argument("--col-b")
    ap.add_argument("--norm", choices=("exact", "space", "case", "nfkc"), default="space")
    ap.add_argument("--sheet-a"); ap.add_argument("--sheet-b")
    a = ap.parse_args()
    if os.path.exists(a.out) and any(os.path.samefile(a.out, p) for p in (a.a, a.b)):
        print("output must differ from both inputs", file=sys.stderr); sys.exit(2)
    norm = normaliser(a.norm)
    lists = {}
    for tag, path, col, sheet in (("A", a.a, a.col, a.sheet_a), ("B", a.b, a.col_b or a.col, a.sheet_b)):
        head, rows = load(path, sheet)
        if col not in head:
            print(f"{path}: column {col!r} not found; header is {head}", file=sys.stderr); sys.exit(2)
        i = head.index(col)
        keys = [norm(r[i]) if i < len(r) else "" for r in rows]
        lists[tag] = dict(path=path, head=head, rows=rows, keys=keys, col=col, idx=i)
    out = {}
    for x, y in (("A", "B"), ("B", "A")):
        L, other = lists[x], set(k for k in lists[y]["keys"] if k)
        missing, blank, matched = [], [], 0
        for n, (r, k) in enumerate(zip(L["rows"], L["keys"]), 2):
            if not k:
                blank.append([x, n] + r)
            elif k in other:
                matched += 1
            else:
                missing.append([n] + r)
        if matched + len(missing) + len(blank) != len(L["rows"]):
            if os.path.exists(a.out):
                os.remove(a.out)
            print(f"SELF-CHECK FAILED for list {x}", file=sys.stderr); sys.exit(3)
        reps = [[x, raw_k, c] for raw_k, c in sorted(Counter(k for k in L["keys"] if k).items()) if c > 1]
        out[x] = dict(missing=missing, blank=blank, matched=matched, reps=reps)
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()

    def sheet(ws, head, rows):
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append(r)
            for c in ws[ws.max_row]:
                if isinstance(c.value, str) and c.value[:1] in ("=", "+", "-", "@"):
                    c.data_type = "s"  # literal text, never a formula
        ws.freeze_panes = "A2"
    sheet(wb.active, ["source_row_in_A"] + lists["A"]["head"], out["A"]["missing"]); wb.active.title = "Missing from B"
    sheet(wb.create_sheet("Missing from A"), ["source_row_in_B"] + lists["B"]["head"], out["B"]["missing"])
    sheet(wb.create_sheet("Blank names"), ["list", "source_row"] + lists["A"]["head"], out["A"]["blank"] + out["B"]["blank"])
    sheet(wb.create_sheet("Repeated names"), ["list", "normalised_name", "rows_with_this_name"], out["A"]["reps"] + out["B"]["reps"])
    note = [
        ["List A", lists["A"]["path"], f"key column {lists['A']['col']!r}", f"{len(lists['A']['rows'])} rows"],
        ["List B", lists["B"]["path"], f"key column {lists['B']['col']!r}", f"{len(lists['B']['rows'])} rows"],
        ["Normalisation", a.norm, normaliser.__doc__ or "", "see the tool header for the exact rule; accents are never removed"],
        ["A: matched / missing from B / blank", out["A"]["matched"], len(out["A"]["missing"]), len(out["A"]["blank"])],
        ["B: matched / missing from A / blank", out["B"]["matched"], len(out["B"]["missing"]), len(out["B"]["blank"])],
        ["Policy", "a name is present if it occurs at least once in the other list; blanks never match; repeated names are listed, not merged", "", ""],
    ]
    sheet(wb.create_sheet("Note"), ["item", "value", "", ""], note)
    wb.save(a.out)
    for x, y in (("A", "B"), ("B", "A")):
        o = out[x]
        print(f"{x}: {len(lists[x]['rows'])} rows = {o['matched']} matched + {len(o['missing'])} missing from {y} + {len(o['blank'])} blank; "
              f"{len(o['reps'])} repeated name(s)")
    print(f"normalisation: {a.norm} -> {a.out}")


if __name__ == "__main__":
    main()
