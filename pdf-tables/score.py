"""Score OUT xlsx against truth json: every truth cell -> written wrong, and wrong AND not yellow (= silent error).
Table extents come from the Summary sheet (rows x cols), so blank-but-yellow rows are not skipped."""
import json, sys
from openpyxl import load_workbook
truth, outdir = json.load(open(sys.argv[1])), sys.argv[2]
tot = wrong = silent = yel = 0
for f, pages in truth.items():
    wb = load_workbook(f"{outdir}/{f[:-4]}.xlsx")
    dims = {(r[0], r[1]): r[2] for r in wb["Summary"].iter_rows(min_row=2, values_only=True)}
    for p, tabs in pages.items():
        ws = wb[f"Page {p}"]; start = 1; got = []
        for k in range(1, len(tabs) + 1):
            d = dims.get((f"Page {p}", k))
            if not d: got.append([]); continue
            nr, nc = map(int, d.split("x"))
            got.append([[ws.cell(start + 1 + i, j + 1) for j in range(nc)] for i in range(nr)]); start += nr + 2
        for T, G in zip(tabs, got):
            for i, tr in enumerate(T):
                for j, tv in enumerate(tr):
                    tot += 1
                    c = G[i][j] if i < len(G) and j < len(G[i]) else None
                    y = c is not None and c.fill.fgColor.rgb in ("00FFFF00", "FFFFFF00"); yel += y
                    if c is None or (c.value or "") != tv:
                        wrong += 1; silent += not y
print(f"truth cells {tot}: written wrong {wrong}, wrong AND not yellow (silent) {silent}, yellow {yel}")
