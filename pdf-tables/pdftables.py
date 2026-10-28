#!/usr/bin/env python3
"""PDF tables (selectable text AND scanned pages) -> one XLSX per PDF, one worksheet per page.

usage: pdftables.py IN_DIR OUT_DIR [--api http://127.0.0.1:8005/v1] [--model NAME] [--dpi 300]

Per page:
  - text layer present -> PyMuPDF find_tables(); cell text is copied from the text layer (no OCR).
  - no text layer (scan) -> deskew, find the ruled grid (horizontal/vertical lines), then read EVERY cell twice:
      reader 1: Tesseract on the cell crop;
      reader 2: a LOCAL vision model reading the whole table crop as a grid of the same shape.
    Cells where the two readers agree are written plainly; cells where they differ are written with reader 1's
    text, filled yellow, and listed on the Exceptions sheet with both readings. Nothing is silently picked.
  - a page with no table gets a one-line note on its worksheet, so every page has a sheet.
Values are written as literal text (no formulas, no number conversion), so "1,200.00" and "007" stay as printed.
The model endpoint must be on this machine unless --allow-remote is given.
"""
import argparse
import base64
import io
import json
import os
import re
import subprocess
import sys
import tempfile
import urllib.parse
import urllib.request

import numpy as np
import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from PIL import Image, ImageFilter

YELLOW = PatternFill("solid", fgColor="FFFF00")


def fit_columns(ws, max_width=60):
    """Readable layout: each column as wide as its longest value (capped), long values wrap; values are unchanged."""
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter
    widths = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is None or (isinstance(c.value, str) and c.value.startswith("Table ") and c.column == 1):
                continue
            widths[c.column] = max(widths.get(c.column, 0), len(str(c.value)))
            if len(str(c.value)) > max_width:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max(w + 2, 8), max_width)


def deskew(g):
    """Angle (deg) that makes horizontal rulings sharpest, searched in +-1.5 deg."""
    small = Image.fromarray(g).resize((g.shape[1] // 3, g.shape[0] // 3))
    best = (0, -1)
    for ang in np.arange(-1.5, 1.51, 0.05):
        a = np.array(small.rotate(ang, fillcolor=255, resample=Image.BILINEAR)) < 128
        score = np.var(a.sum(1))
        if score > best[1]:
            best = (ang, score)
    return best[0]


def lines(mask, axis, min_frac):
    """Positions of ruling lines: rows (axis=1) or columns (axis=0) whose dark share exceeds min_frac."""
    prof = mask.mean(axis)
    idx = np.where(prof > min_frac)[0]
    groups, cur = [], []
    for i in idx:
        if cur and i - cur[-1] > 3:
            groups.append(cur); cur = []
        cur.append(i)
    if cur:
        groups.append(cur)
    return [int(np.mean(g)) for g in groups]


def otsu(g):
    h = np.bincount(g.ravel(), minlength=256).astype(float); w = np.cumsum(h); m = np.cumsum(h * np.arange(256))
    with np.errstate(divide="ignore", invalid="ignore"):
        v = (m[-1] * w - m * w[-1]) ** 2 / (w * (w[-1] - w))
    return int(np.nanargmax(v))


def grids(g):
    """Split a page into ruled tables; returns [(ys, xs)] in page order. Noise is median-filtered first."""
    d = np.array(Image.fromarray(g).filter(ImageFilter.MedianFilter(5)))
    mask = d < otsu(d)
    hs = lines(mask, 1, 0.25)
    if len(hs) < 2:
        return []
    gaps = np.diff(hs)
    step = np.median(gaps)
    tables, cur = [], [hs[0]]
    for y, d in zip(hs[1:], gaps):
        if d > 2.5 * step:  # a large vertical gap starts a new table
            tables.append(cur); cur = []
        cur.append(y)
    tables.append(cur)
    out = []
    for ys in tables:
        if len(ys) < 2:
            continue
        band = mask[ys[0]:ys[-1] + 1]
        xs = lines(band, 0, 0.7)
        if len(xs) >= 2:
            out.append((ys, xs))
    return out


def ocr_cell(img):
    with tempfile.NamedTemporaryFile(suffix=".png") as f:
        img.save(f.name)
        r = subprocess.run(["tesseract", f.name, "-", "--psm", "7", "-l", "eng"], capture_output=True, text=True)
    return re.sub(r"\s+", " ", r.stdout).strip()


def ask_grid(api, model, img):
    buf = io.BytesIO(); img.save(buf, "PNG")
    # the grid shape is NOT given to the model, so its row/column count is an independent check on grid detection
    prompt = ("This image is a ruled table. Transcribe every cell exactly as printed, including punctuation, commas "
              'and decimals, row by row. Return JSON only: {"rows": [[cell, ...], ...]}; use "" for an empty cell.')
    body = {"model": model, "temperature": 0, "response_format": {"type": "json_object"},
            "messages": [{"role": "user", "content": [
                {"type": "text", "text": prompt},
                {"type": "image_url", "image_url": {"url": "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()}}]}]}
    req = urllib.request.Request(api + "/chat/completions", json.dumps(body).encode(), {"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=600) as r:
        out = json.load(r)["choices"][0]["message"]["content"]
    return json.loads(re.sub(r"^```(?:json)?|```$", "", out.strip(), flags=re.M))["rows"]


def norm(s):
    return re.sub(r"\s+", " ", str(s)).strip()


def scanned_tables(page, a, exc, pname):
    pix = page.get_pixmap(dpi=a.dpi, colorspace=pymupdf.csGRAY)
    im = Image.frombytes("L", (pix.width, pix.height), pix.samples)
    ang = deskew(np.array(im))
    im = im.rotate(ang, fillcolor=255, resample=Image.BICUBIC)
    g = np.array(im)
    result = []
    for t, (ys, xs) in enumerate(grids(g), 1):
        nr, nc = len(ys) - 1, len(xs) - 1
        pad = 4
        r1 = [[ocr_cell(im.crop((xs[j] + pad, ys[i] + pad, xs[j + 1] - pad, ys[i + 1] - pad))) for j in range(nc)] for i in range(nr)]
        crop = im.crop((xs[0] - 6, ys[0] - 6, xs[-1] + 6, ys[-1] + 6))
        try:
            r2 = ask_grid(a.api, a.model, crop)
        except Exception as e:  # reader 2 unavailable -> every cell is unconfirmed
            r2 = None; exc.append([pname, t, "", "", f"reader 2 failed ({e}); whole table unconfirmed", "", ""])
        if r2 is not None and (len(r2) != nr or any(len(r) != nc for r in r2)):
            exc.append([pname, t, "", "", f"reader 2 read {len(r2)} rows x {max(len(r) for r in r2) if r2 else 0} cols, ruled grid is {nr}x{nc}; grid or read is wrong - whole table unconfirmed", "", ""])
            r2 = None
        flags = []
        for i in range(nr):
            for j in range(nc):
                v2 = norm(r2[i][j]) if r2 is not None else None
                if v2 != r1[i][j]:
                    flags.append((i, j))
                    if r2 is not None:
                        exc.append([pname, t, i + 1, j + 1, "readers differ", r1[i][j], v2])
        if len(flags) > 0.2 * nr * nc:  # too many disagreements: the scan is below the agreed quality, not "mostly right"
            exc.append([pname, t, "", "", f"{len(flags)} of {nr * nc} cells unconfirmed - scan below eligibility; re-scan or check by hand", "", ""])
            flags = [(i, j) for i in range(nr) for j in range(nc)]
        result.append(dict(rows=r1, flags=flags, how=f"scan, deskew {ang:+.2f} deg, {nr}x{nc} ruled grid, two readers"))
    return result


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("inp"); ap.add_argument("out")
    ap.add_argument("--api", default="http://127.0.0.1:8005/v1"); ap.add_argument("--model", default="local")
    ap.add_argument("--dpi", type=int, default=300); ap.add_argument("--allow-remote", action="store_true")
    a = ap.parse_args()
    if urllib.parse.urlparse(a.api).hostname not in ("127.0.0.1", "localhost", "::1") and not a.allow_remote:
        sys.exit("refusing a non-local model endpoint (customer files stay on this machine)")
    os.makedirs(a.out, exist_ok=True)
    for name in sorted(f for f in os.listdir(a.inp) if f.lower().endswith(".pdf")):
        doc = pymupdf.open(os.path.join(a.inp, name))
        wb = Workbook(); wb.remove(wb.active); exc = []; summary = []
        for page in doc:
            pname = f"Page {page.number + 1}"
            ws = wb.create_sheet(pname)
            if len(page.get_text().strip()) >= 20:
                tabs = [dict(rows=[[norm(c or "") for c in r] for r in t.extract()], flags=[], how="text layer")
                        for t in page.find_tables().tables]
            else:
                tabs = scanned_tables(page, a, exc, pname)
            r = 1
            if not tabs:
                ws.cell(1, 1, "No table found on this page."); summary.append([pname, 0, "", 0])
                if len(page.get_text().strip()) < 20:  # a scan where no ruled grid was found is never silently "no table"
                    exc.append([pname, "", "", "", "scanned page: no ruled table detected - check this page by hand", "", ""])
            for k, t in enumerate(tabs, 1):
                ws.cell(r, 1, f"Table {k} ({t['how']})").font = Font(italic=True); r += 1
                for i, row in enumerate(t["rows"]):
                    for j, v in enumerate(row):
                        c = ws.cell(r + i, j + 1, v); c.data_type = "s"
                        if (i, j) in t["flags"]:
                            c.fill = YELLOW
                r += len(t["rows"]) + 1
                summary.append([pname, k, f"{len(t['rows'])}x{len(t['rows'][0]) if t['rows'] else 0}", len(t["flags"])])
        ws = wb.create_sheet("Summary")
        for row in [["page", "table", "rows x cols", "unconfirmed cells (yellow)"]] + summary:
            ws.append(row)
        ws = wb.create_sheet("Exceptions")
        for row in [["page", "table", "row", "col", "issue", "reader 1 (written)", "reader 2"]] + exc:
            ws.append(row)
        for ws in wb.worksheets:
            fit_columns(ws)
        out = os.path.join(a.out, os.path.splitext(name)[0] + ".xlsx")
        wb.save(out)
        print(f"{name}: {len(doc)} pages -> {out}; {sum(s[3] for s in summary)} unconfirmed cells, {len(exc)} exceptions")


if __name__ == "__main__":
    main()
