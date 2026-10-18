#!/usr/bin/env python3
"""Plain-text data file -> cleaned XLSX: whole-row dedup, every blank filled with one placeholder, no empty cells.

usage: txt2xlsx.py IN.txt OUT.xlsx --placeholder "N/A" [--delimiter auto|,|;|\\t||] [--encoding utf-8]
                   [--order fill-then-dedup(default)|dedup-then-fill] [--normalize-ws] [--null-markers "NULL,NA,-"]
                   [--numeric-cols "amount"]
Never modifies the input. Writes OUT.xlsx (sheet "Data" + sheet "Log") and prints a reconciliation summary.
"""
import argparse
import csv
import io
import os
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


def sniff(text, forced):
    if forced != "auto":
        return "\t" if forced in ("\\t", "tab") else forced
    # pick the candidate whose per-line field count is most consistent (and > 1) over the first lines
    lines = [ln for ln in text.splitlines()[:200] if ln.strip()]
    best, score = ",", -1
    for d in (",", ";", "\t", "|"):
        counts = [len(next(csv.reader([ln], delimiter=d))) for ln in lines]
        common = max(set(counts), key=counts.count)
        sc = counts.count(common) if common > 1 else 0
        if sc > score:
            best, score = d, sc
    return best


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("dst")
    ap.add_argument("--placeholder", required=True)
    ap.add_argument("--delimiter", default="auto")
    ap.add_argument("--encoding", default="utf-8-sig")
    ap.add_argument("--order", choices=["dedup-then-fill", "fill-then-dedup"], default="fill-then-dedup")  # matches a COUNTBLANK=0 + no-duplicates acceptance check
    ap.add_argument("--normalize-ws", action="store_true", help="trim cells and treat whitespace-only as missing")
    ap.add_argument("--null-markers", default="", help="comma list of text treated as missing, e.g. NULL,NA")
    ap.add_argument("--numeric-cols", default="",
                    help="columns written as numbers (default: every cell is kept as its exact source text)")
    a = ap.parse_args()
    if os.path.exists(a.dst) and (os.path.realpath(a.src) == os.path.realpath(a.dst) or os.path.samefile(a.src, a.dst)):
        print("output must differ from input (same file, alias or hard link)", file=sys.stderr)
        return 2
    if os.path.realpath(a.src) == os.path.realpath(a.dst):
        print("output must differ from input", file=sys.stderr)
        return 2

    text = open(a.src, encoding=a.encoding, newline="").read()
    delim = sniff(text, a.delimiter)
    rows = list(csv.reader(io.StringIO(text), delimiter=delim))
    header, body = rows[0], rows[1:]
    width = max([len(header)] + [len(r) for r in body])
    if len(header) < width:
        header += [f"column_{i + 1}" for i in range(len(header), width)]
    nulls = {m.strip() for m in a.null_markers.split(",") if m.strip()}
    log = [("input rows (excl. header)", len(body)), ("delimiter", repr(delim)), ("columns", width)]

    short = sum(1 for r in body if len(r) < width)
    body = [r + [""] * (width - len(r)) for r in body]  # ragged rows padded -> counted as missing cells
    log.append(("rows padded to full width", short))

    def clean(v):
        if a.normalize_ws:
            v = v.strip()
        return "" if v in nulls else v

    body = [[clean(v) for v in r] for r in body]
    blank_rows = sum(1 for r in body if not any(r))
    body = [r for r in body if any(r)]
    log.append(("fully empty rows removed", blank_rows))

    num_cols = {c.strip() for c in a.numeric_cols.split(",") if c.strip()}

    def typed(col, v):
        """Value as it will be stored: exact text unless the column is declared numeric and v parses."""
        if col in num_cols and v != a.placeholder:
            try:  # 1, 1.0 and 1.00 are the same number in Excel: normalise so dedup sees what Excel sees
                f = float(v)
                return int(f) if f.is_integer() and abs(f) < 2**53 else f
            except ValueError:
                pass
        return v

    def dedup(rs):  # equality on the values exactly as they will be written (type-aware)
        seen, out = set(), []
        for r in rs:
            k = tuple((type(x).__name__, x) for x in (typed(c, v) for c, v in zip(header, r)))
            if k not in seen:
                seen.add(k)
                out.append(r)
        return out, len(rs) - len(out)

    def fill(rs):
        n = sum(v == "" for r in rs for v in r)
        return [[v if v != "" else a.placeholder for v in r] for r in rs], n

    if a.order == "dedup-then-fill":
        body, dups = dedup(body)
        body, filled = fill(body)
    else:
        body, filled = fill(body)
        body, dups = dedup(body)
    log += [("order", a.order), ("equality", "exact source text" if not num_cols else
             f"exact text; numeric compare in {sorted(num_cols)}"),
            ("duplicate rows removed (first kept, original order)", dups),
            ("missing cells filled with placeholder", filled), ("output rows", len(body))]
    if a.order == "dedup-then-fill":  # rows that only became identical after filling (kept, flagged)
        _, extra = dedup(body)
        log.append(("rows identical only after filling (kept; agree policy)", extra))

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"

    def put(r_idx, c_idx, value, bold=False):
        cell = ws.cell(row=r_idx, column=c_idx)
        cell.value = value
        if isinstance(value, str):  # force literal text: "=1+1" must never become a formula
            cell.data_type = "s"
            cell.number_format = "@"
        if bold:
            cell.font = Font(bold=True)

    for j, h in enumerate(header, 1):
        put(1, j, h, bold=True)
    for i, r in enumerate(body, 2):
        for j, (col, v) in enumerate(zip(header, r), 1):
            put(i, j, typed(col, v))
    lg = wb.create_sheet("Log")
    lg.append(["step", "value"])
    for k, v in log:
        lg.append([k, str(v)])
    wb.save(a.dst)

    # verify the written file as stored values: no formulas, no empty cells, no duplicate rows
    raw = load_workbook(a.dst)["Data"]
    formulas = sum(c.data_type == "f" for row in raw.iter_rows(max_col=width) for c in row)
    vals = [[c.value for c in row] for row in load_workbook(a.dst, data_only=True)["Data"].iter_rows(min_row=2, max_col=width)]
    empty = sum(v is None or v == "" for r in vals for v in r)
    dup_rows = len(vals) - len({tuple((type(x).__name__, x) for x in r) for r in vals})
    print("\n".join(f"{k}: {v}" for k, v in log))
    print(f"VERIFY data range A1:{raw.cell(1, width).column_letter}{len(vals) + 1}: formula cells={formulas}, "
          f"empty cells={empty}, duplicate rows={dup_rows}")
    return 0 if formulas == 0 and empty == 0 and dup_rows == 0 else 1

if __name__ == "__main__":
    sys.exit(main())
